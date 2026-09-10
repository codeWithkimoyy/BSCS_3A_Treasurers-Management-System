import csv
from datetime import date, datetime
import io

from flask import Blueprint, flash, redirect, render_template, request, send_file, url_for
from flask_login import login_required
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .audit import log_audit
from .db import db, next_id
from .decorators import current_user_id, get_float, role_required, validate_csrf

bp = Blueprint('events', __name__)

ROLES_MANAGE = ('admin', 'mayor', 'treasurer')


@bp.route('/events')
@login_required
def index():
    filter_status = request.args.get('status', 'active')

    query = {}
    if filter_status == 'active':
        query['status'] = 'active'
    elif filter_status == 'closed':
        query['status'] = 'closed'

    evts = list(db.events.find(query).sort('created_at', -1))
    total_students = db.students.count_documents({'is_active': 1})
    active_count = db.events.count_documents({'status': 'active'})
    closed_count = db.events.count_documents({'status': 'closed'})
    all_count = db.events.count_documents({})

    for e in evts:
        e['amount'] = float(e.get('amount', 0) or 0)
        total_paid = list(db.payments.aggregate([
            {'$match': {'event_id': e['_id'], 'confirmed': True}},
            {'$group': {'_id': None, 'total': {'$sum': '$amount_paid'}}}
        ]))
        paid_count = db.payments.count_documents({'event_id': e['_id'], 'confirmed': True})
        locked_count = db.payments.count_documents({'event_id': e['_id'], 'locked': True})
        e['collected'] = float(total_paid[0]['total'] if total_paid else 0)
        e['paid_count'] = paid_count
        e['total_students'] = total_students
        e['fully_paid'] = locked_count == total_students if total_students > 0 else False

    return render_template('events.html', events=evts, filter_status=filter_status,
        active_count=active_count, closed_count=closed_count, all_count=all_count)


@bp.route('/events/add', methods=['POST'])
@login_required
@role_required(*ROLES_MANAGE)
def add():
    validate_csrf()
    amount = get_float(request.form, 'amount')
    db.events.insert_one({
        '_id': next_id('events'),
        'title': request.form['title'].strip(),
        'amount': amount,
        'deadline': request.form.get('deadline', ''),
        'status': 'active',
        'created_by': current_user_id(),
        'created_at': datetime.now(),
    })
    log_audit('event_add', target=request.form['title'])
    flash('Event created', 'success')
    return redirect(url_for('events.index'))


@bp.route('/events/edit/<int:id>', methods=['POST'])
@login_required
@role_required(*ROLES_MANAGE)
def edit(id):
    validate_csrf()
    amount = get_float(request.form, 'amount')
    db.events.update_one({'_id': id}, {'$set': {
        'title': request.form['title'].strip(),
        'amount': amount,
        'deadline': request.form.get('deadline', ''),
    }})
    log_audit('event_edit', target=id)
    flash('Event updated', 'success')
    return redirect(url_for('events.index'))


@bp.route('/events/close/<int:id>', methods=['POST'])
@login_required
@role_required(*ROLES_MANAGE)
def close(id):
    validate_csrf()
    db.events.update_one({'_id': id}, {'$set': {'status': 'closed'}})
    log_audit('event_close', target=id)
    flash('Event closed', 'success')
    return redirect(url_for('events.index'))


@bp.route('/events/reopen/<int:id>', methods=['POST'])
@login_required
@role_required(*ROLES_MANAGE)
def reopen(id):
    validate_csrf()
    db.events.update_one({'_id': id}, {'$set': {'status': 'active'}})
    log_audit('event_reopen', target=id)
    flash('Event reopened for active collections', 'success')
    return redirect(url_for('events.index'))


@bp.route('/events/finalize/<int:id>', methods=['POST'])
@login_required
@role_required(*ROLES_MANAGE)
def finalize(id):
    validate_csrf()
    event = db.events.find_one({'_id': id})
    if not event:
        flash('Event not found', 'error')
        return redirect(url_for('events.index'))
    allocation = request.form.get('allocation', '')
    if allocation not in ('balance', 'expense'):
        flash('Invalid allocation', 'error')
        return redirect(url_for('events.index'))
    total = list(db.payments.aggregate([
        {'$match': {'event_id': id, 'locked': True}},
        {'$group': {'_id': None, 'total': {'$sum': '$amount_paid'}}}
    ]))
    collected = total[0]['total'] if total else 0
    if allocation == 'expense' and collected > 0:
        db.transactions.insert_one({
            '_id': next_id('transactions'),
            'type': 'expense',
            'amount': collected,
            'description': f"Event allocation: {event.get('title', '')}",
            'transaction_date': date.today().isoformat(),
            'created_by': current_user_id(),
            'created_at': datetime.now(),
        })
    db.events.update_one({'_id': id}, {'$set': {
        'status': 'closed',
        'allocation': allocation,
        'allocated_at': datetime.now(),
        'allocated_by': current_user_id(),
    }})
    log_audit('event_finalize', target=id, detail=allocation)
    flash(f'Event finalized — funds {"kept in balance" if allocation == "balance" else "recorded as expense"}', 'success')
    return redirect(url_for('events.index'))


@bp.route('/events/delete/<int:id>', methods=['POST'])
@login_required
@role_required(*ROLES_MANAGE)
def delete(id):
    validate_csrf()
    db.payments.delete_many({'event_id': id})
    db.events.delete_one({'_id': id})
    log_audit('event_delete', target=id)
    flash('Event deleted', 'success')
    return redirect(url_for('events.index'))


def get_event_payment_roster(event_id):
    event = db.events.find_one({'_id': event_id})
    if not event:
        return None, None
    students = list(db.students.find({'is_active': 1, 'deleted': {'$ne': True}}).sort('name', 1))
    payments = {p['student_id']: p for p in db.payments.find({'event_id': event_id})}
    user_map = {u['_id']: u.get('display_name') or u.get('username') for u in db.users.find({})}

    target = float(event.get('amount', 0))
    roster = []
    for idx, s in enumerate(students, 1):
        pmt = payments.get(s['_id'])
        paid = float(pmt.get('amount_paid', 0)) if pmt else 0.0
        balance = max(0.0, target - paid)
        if pmt and pmt.get('locked'):
            status = 'Fully Paid'
        elif paid > 0:
            status = 'Partial'
        else:
            status = 'Pending'

        conf_by = user_map.get(pmt.get('confirmed_by'), '') if pmt else ''
        conf_at = pmt.get('confirmed_at').strftime('%Y-%m-%d %H:%M') if pmt and pmt.get('confirmed_at') else ''

        roster.append({
            'num': idx,
            'student_id': s.get('student_id', ''),
            'name': s.get('name', ''),
            'course': s.get('course', ''),
            'year': s.get('year', ''),
            'target': target,
            'paid': paid,
            'balance': balance,
            'status': status,
            'confirmed_at': conf_at,
            'confirmed_by': conf_by,
            'notes': pmt.get('notes', '') if pmt else ''
        })
    return event, roster


@bp.route('/events/export/<int:id>')
@login_required
def export_csv(id):
    event, roster = get_event_payment_roster(id)
    if not event:
        flash('Event not found', 'error')
        return redirect(url_for('events.index'))

    output = io.StringIO()
    w = csv.writer(output)
    w.writerow(['EVENT PAYMENT ROSTER & AUDIT REPORT'])
    w.writerow(['Event Title', event.get('title', '')])
    w.writerow(['Fee Amount (PHP)', f"{event.get('amount', 0):.2f}"])
    w.writerow(['Deadline', event.get('deadline', 'N/A')])
    w.writerow(['Total Members', len(roster)])
    total_collected = sum(r['paid'] for r in roster)
    total_balance = sum(r['balance'] for r in roster)
    w.writerow(['Total Collected (PHP)', f"{total_collected:.2f}"])
    w.writerow(['Total Outstanding (PHP)', f"{total_balance:.2f}"])
    w.writerow(['Generated At', datetime.now().strftime('%Y-%m-%d %H:%M:%S')])
    w.writerow([])
    w.writerow(['No.', 'Student ID', 'Full Name', 'Course', 'Year', 'Target (PHP)', 'Amount Paid (PHP)', 'Balance Due (PHP)', 'Status', 'Date Confirmed', 'Confirmed By', 'Notes'])

    for r in roster:
        w.writerow([
            r['num'], r['student_id'], r['name'], r['course'], r['year'],
            f"{r['target']:.2f}", f"{r['paid']:.2f}", f"{r['balance']:.2f}",
            r['status'], r['confirmed_at'], r['confirmed_by'], r['notes']
        ])

    mem = io.BytesIO()
    mem.write(output.getvalue().encode('utf-8-sig'))
    mem.seek(0)
    slug = "".join(c if c.isalnum() else "_" for c in event.get('title', 'event'))
    return send_file(mem, as_attachment=True, download_name=f'event_{slug}_{date.today().isoformat()}.csv', mimetype='text/csv')


@bp.route('/events/export_excel/<int:id>')
@login_required
def export_excel(id):
    event, roster = get_event_payment_roster(id)
    if not event:
        flash('Event not found', 'error')
        return redirect(url_for('events.index'))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Event Payments"[:31]

    navy_fill = PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    title_font = Font(name="Arial", size=13, bold=True, color="1E1B4B")
    sub_font = Font(name="Arial", size=9, italic=True, color="64748B")
    bold_font = Font(name="Arial", size=10, bold=True)
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    ws['A1'] = "OFFICIAL EVENT PAYMENT ROSTER"
    ws['A1'].font = title_font
    ws['A2'] = f"Event: {event.get('title', '')} | Fee: ₱{event.get('amount', 0):.2f} | Deadline: {event.get('deadline') or 'N/A'}"
    ws['A2'].font = sub_font
    ws['A3'] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ws['A3'].font = sub_font

    headers = ['No.', 'Student ID', 'Full Name', 'Course', 'Year', 'Target (₱)', 'Amount Paid (₱)', 'Balance (₱)', 'Status', 'Confirmed Date', 'Confirmed By']
    start_row = 5
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=h)
        cell.fill = navy_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center" if col_idx in (1, 4, 5, 9) else ("right" if col_idx in (6, 7, 8) else "left"))

    for idx, r in enumerate(roster, start_row + 1):
        ws.cell(row=idx, column=1, value=r['num']).alignment = Alignment(horizontal="center")
        ws.cell(row=idx, column=2, value=r['student_id']).alignment = Alignment(horizontal="center")
        ws.cell(row=idx, column=3, value=r['name'])
        ws.cell(row=idx, column=4, value=r['course']).alignment = Alignment(horizontal="center")
        ws.cell(row=idx, column=5, value=r['year']).alignment = Alignment(horizontal="center")

        c_target = ws.cell(row=idx, column=6, value=r['target'])
        c_target.number_format = '₱#,##0.00'
        c_paid = ws.cell(row=idx, column=7, value=r['paid'])
        c_paid.number_format = '₱#,##0.00'
        c_bal = ws.cell(row=idx, column=8, value=r['balance'])
        c_bal.number_format = '₱#,##0.00'

        c_status = ws.cell(row=idx, column=9, value=r['status'])
        c_status.alignment = Alignment(horizontal="center")
        if r['status'] == 'Fully Paid':
            c_status.font = Font(name="Arial", size=10, bold=True, color="059669")
        elif r['status'] == 'Partial':
            c_status.font = Font(name="Arial", size=10, bold=True, color="D97706")
        else:
            c_status.font = Font(name="Arial", size=10, color="94A3B8")

        ws.cell(row=idx, column=10, value=r['confirmed_at'])
        ws.cell(row=idx, column=11, value=r['confirmed_by'])

        for c in range(1, 12):
            ws.cell(row=idx, column=c).border = thin_border

    last_row = start_row + len(roster)
    total_row = last_row + 1
    ws.cell(row=total_row, column=3, value="TOTALS").font = bold_font
    ws.cell(row=total_row, column=6, value=f"=SUM(F{start_row+1}:F{last_row})").number_format = '₱#,##0.00'
    ws.cell(row=total_row, column=7, value=f"=SUM(G{start_row+1}:G{last_row})").number_format = '₱#,##0.00'
    ws.cell(row=total_row, column=8, value=f"=SUM(H{start_row+1}:H{last_row})").number_format = '₱#,##0.00'
    for c in (3, 6, 7, 8):
        ws.cell(row=total_row, column=c).font = bold_font

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    mem = io.BytesIO()
    wb.save(mem)
    mem.seek(0)
    slug = "".join(c if c.isalnum() else "_" for c in event.get('title', 'event'))
    return send_file(mem, as_attachment=True, download_name=f'event_{slug}_{date.today().isoformat()}.xlsx', mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@bp.route('/events/print/<int:id>')
@login_required
def print_roster(id):
    event, roster = get_event_payment_roster(id)
    if not event:
        flash('Event not found', 'error')
        return redirect(url_for('events.index'))

    total_collected = sum(r['paid'] for r in roster)
    total_balance = sum(r['balance'] for r in roster)
    total_target = sum(r['target'] for r in roster)
    paid_count = sum(1 for r in roster if r['status'] == 'Fully Paid')

    return render_template(
        'event_print.html',
        event=event,
        roster=roster,
        total_collected=total_collected,
        total_balance=total_balance,
        total_target=total_target,
        paid_count=paid_count,
        now_str=datetime.now().strftime('%B %d, %Y %I:%M %p')
    )


@bp.route('/events/export_all_excel')
@login_required
def export_all_excel():
    events = list(db.events.find({'deleted': {'$ne': True}}).sort('created_at', 1))
    students = list(db.students.find({'is_active': 1, 'deleted': {'$ne': True}}).sort('name', 1))
    user_map = {u['_id']: u.get('display_name') or u.get('username') for u in db.users.find({})}

    wb = openpyxl.Workbook()
    ws_ov = wb.active
    ws_ov.title = "Overview"

    navy_fill = PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    bold_font = Font(name="Arial", size=10, bold=True)
    title_font = Font(name="Arial", size=13, bold=True, color="1E1B4B")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    ws_ov['A1'] = "CLASS TREASURY - ALL EVENTS SUMMARY"
    ws_ov['A1'].font = title_font
    ws_ov['A2'] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ws_ov['A2'].font = Font(name="Arial", size=9, italic=True, color="64748B")

    ov_headers = ['#', 'Event Title', 'Deadline', 'Status', 'Fee (₱)', 'Eligible Students', 'Expected Total (₱)', 'Collected (₱)', 'Balance (₱)', 'Completion']
    for c_idx, h in enumerate(ov_headers, 1):
        cell = ws_ov.cell(row=4, column=c_idx, value=h)
        cell.fill = navy_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center" if c_idx in (1, 3, 4, 6, 10) else ("right" if c_idx in (5, 7, 8, 9) else "left"))

    used_sheet_names = {"Overview"}
    for ev_idx, ev in enumerate(events, 1):
        ev_id = ev['_id']
        ev_title = ev.get('title', f'Event {ev_idx}')
        fee = float(ev.get('amount', 0))
        total_students = len(students)
        expected = fee * total_students

        pmts = {p['student_id']: p for p in db.payments.find({'event_id': ev_id})}
        collected = sum(float(p.get('amount_paid', 0)) for p in pmts.values() if p.get('confirmed'))
        balance = max(0.0, expected - collected)
        pct = (collected / expected * 100) if expected > 0 else 0

        row = 4 + ev_idx
        ws_ov.cell(row=row, column=1, value=ev_idx).alignment = Alignment(horizontal="center")
        ws_ov.cell(row=row, column=2, value=ev_title)
        ws_ov.cell(row=row, column=3, value=ev.get('deadline') or '—').alignment = Alignment(horizontal="center")
        ws_ov.cell(row=row, column=4, value=ev.get('status', 'active').title()).alignment = Alignment(horizontal="center")

        c_fee = ws_ov.cell(row=row, column=5, value=fee)
        c_fee.number_format = '₱#,##0.00'
        ws_ov.cell(row=row, column=6, value=total_students).alignment = Alignment(horizontal="center")

        c_exp = ws_ov.cell(row=row, column=7, value=expected)
        c_exp.number_format = '₱#,##0.00'
        c_col = ws_ov.cell(row=row, column=8, value=collected)
        c_col.number_format = '₱#,##0.00'
        c_bal = ws_ov.cell(row=row, column=9, value=balance)
        c_bal.number_format = '₱#,##0.00'
        ws_ov.cell(row=row, column=10, value=f"{pct:.1f}%").alignment = Alignment(horizontal="center")

        for c in range(1, 11):
            ws_ov.cell(row=row, column=c).border = thin_border

        safe_name = "".join(c for c in ev_title if c.isalnum() or c in (' ', '_', '-'))[:28].strip() or f"Event_{ev_idx}"
        sheet_name = safe_name
        counter = 1
        while sheet_name.lower() in [s.lower() for s in used_sheet_names]:
            sheet_name = f"{safe_name[:25]}_{counter}"
            counter += 1
        used_sheet_names.add(sheet_name)

        ws_ev = wb.create_sheet(title=sheet_name)
        ws_ev['A1'] = f"{ev_title.upper()} - PAYMENT LEDGER"
        ws_ev['A1'].font = title_font
        ws_ev['A2'] = f"Fee: ₱{fee:.2f} | Deadline: {ev.get('deadline') or 'N/A'}"
        ws_ev['A2'].font = Font(name="Arial", size=9, italic=True, color="64748B")

        ev_headers = ['No.', 'Student ID', 'Full Name', 'Course', 'Year', 'Target (₱)', 'Amount Paid (₱)', 'Balance (₱)', 'Status', 'Confirmed Date', 'Confirmed By']
        for c_idx, h in enumerate(ev_headers, 1):
            c = ws_ev.cell(row=4, column=c_idx, value=h)
            c.fill = navy_fill
            c.font = header_font
            c.alignment = Alignment(horizontal="center" if c_idx in (1, 4, 5, 9) else ("right" if c_idx in (6, 7, 8) else "left"))

        for s_idx, s in enumerate(students, 1):
            p = pmts.get(s['_id'])
            p_paid = float(p.get('amount_paid', 0)) if p else 0.0
            p_bal = max(0.0, fee - p_paid)
            if p and p.get('locked'):
                p_status = 'Fully Paid'
            elif p_paid > 0:
                p_status = 'Partial'
            else:
                p_status = 'Pending'

            r_idx = 4 + s_idx
            ws_ev.cell(row=r_idx, column=1, value=s_idx).alignment = Alignment(horizontal="center")
            ws_ev.cell(row=r_idx, column=2, value=s.get('student_id', '')).alignment = Alignment(horizontal="center")
            ws_ev.cell(row=r_idx, column=3, value=s.get('name', ''))
            ws_ev.cell(row=r_idx, column=4, value=s.get('course', '')).alignment = Alignment(horizontal="center")
            ws_ev.cell(row=r_idx, column=5, value=s.get('year', '')).alignment = Alignment(horizontal="center")

            ws_ev.cell(row=r_idx, column=6, value=fee).number_format = '₱#,##0.00'
            ws_ev.cell(row=r_idx, column=7, value=p_paid).number_format = '₱#,##0.00'
            ws_ev.cell(row=r_idx, column=8, value=p_bal).number_format = '₱#,##0.00'

            c_stat = ws_ev.cell(row=r_idx, column=9, value=p_status)
            c_stat.alignment = Alignment(horizontal="center")
            if p_status == 'Fully Paid':
                c_stat.font = Font(name="Arial", size=10, bold=True, color="059669")
            elif p_status == 'Partial':
                c_stat.font = Font(name="Arial", size=10, bold=True, color="D97706")
            else:
                c_stat.font = Font(name="Arial", size=10, color="94A3B8")

            c_date = p.get('confirmed_at').strftime('%Y-%m-%d %H:%M') if p and p.get('confirmed_at') else ''
            c_user = user_map.get(p.get('confirmed_by'), '') if p else ''
            ws_ev.cell(row=r_idx, column=10, value=c_date)
            ws_ev.cell(row=r_idx, column=11, value=c_user)

            for c in range(1, 12):
                ws_ev.cell(row=r_idx, column=c).border = thin_border

        tot_row = 5 + len(students)
        ws_ev.cell(row=tot_row, column=3, value="TOTALS").font = bold_font
        ws_ev.cell(row=tot_row, column=6, value=f"=SUM(F5:F{tot_row-1})").number_format = '₱#,##0.00'
        ws_ev.cell(row=tot_row, column=7, value=f"=SUM(G5:G{tot_row-1})").number_format = '₱#,##0.00'
        ws_ev.cell(row=tot_row, column=8, value=f"=SUM(H5:H{tot_row-1})").number_format = '₱#,##0.00'
        for c in (3, 6, 7, 8):
            ws_ev.cell(row=tot_row, column=c).font = bold_font

        for col in ws_ev.columns:
            m_len = max(len(str(cell.value or '')) for cell in col)
            ws_ev.column_dimensions[get_column_letter(col[0].column)].width = max(m_len + 3, 11)

    for col in ws_ov.columns:
        m_len = max(len(str(cell.value or '')) for cell in col)
        ws_ov.column_dimensions[get_column_letter(col[0].column)].width = max(m_len + 3, 11)

    mem = io.BytesIO()
    wb.save(mem)
    mem.seek(0)
    return send_file(mem, as_attachment=True, download_name=f'all_events_treasury_{date.today().isoformat()}.xlsx', mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
